"""
日計報告書ExcelファイルをパースしてDictで返すモジュール
"""
import openpyxl


def clean_number(val):
    """セルの値から¥や,を除去して数値に変換。変換できない場合はNoneを返す"""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace('¥', '').replace(',', '').replace(' ', '').strip()
    if s in ('', '-', '¥-', '0%'):
        return None
    # %を除去して割合として返す
    if s.endswith('%'):
        try:
            return float(s[:-1])
        except ValueError:
            return None
    try:
        return float(s)
    except ValueError:
        return None


def find_col(ws, header_row_idx, keywords):
    """
    ヘッダー行（2行分）を検索して、最初にkeywordsのいずれかに一致した列番号(1-indexed)を返す
    """
    for row_offset in (0, 1):
        row_idx = header_row_idx + row_offset
        if row_idx > ws.max_row:
            break
        for cell in ws[row_idx]:
            if not cell.value:
                continue
            cell_str = str(cell.value).replace('\n', '')
            for kw in keywords:
                if kw in cell_str:
                    return cell.column
    return None


def parse_daily_report(filepath):
    """
    日計報告書ExcelファイルをパースしてDictで返す

    Returns:
        {
            'store': '高円寺',   # 店舗名（「店」を除いた形）
            'stylists': [
                {
                    'name':    '望月',
                    '生産性':  4693.0,
                    '売上':    774420.0,
                    '稼働時間': 150.0,
                    '客単価':  9928.0,
                    '新規数':  2.0,
                    '指名率':  70.51,   # % 表記を除いた数値
                },
                ...
            ]
        }
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    # ── 1行目から店舗名を取得 ──────────────────────────────
    store_name = None
    for cell in ws[1]:
        val = str(cell.value or '').strip()
        if '店' in val and len(val) > 1:
            store_name = val.replace('店', '').strip()
            break

    # ── ヘッダー行を探す（「スタイリスト」セルがある行） ────
    header_row_idx = None
    stylist_col = None

    for row in ws.iter_rows(max_row=30):
        for cell in row:
            if cell.value and 'スタイリスト' in str(cell.value):
                header_row_idx = cell.row
                stylist_col = cell.column
                break
        if header_row_idx:
            break

    if not header_row_idx:
        raise ValueError(f"ヘッダー行が見つかりません: {filepath}")

    # ── 各指標の列インデックスを取得 ──────────────────────
    col_生産性  = find_col(ws, header_row_idx, ['生産性（税抜）', '生産性'])
    col_売上   = find_col(ws, header_row_idx, ['合計売上(税込)', '合計売上'])
    col_稼働   = find_col(ws, header_row_idx, ['勤務時間（15分', '勤務時間\n（15分', '勤務時間'])
    col_客単価  = find_col(ws, header_row_idx, ['客単価'])
    col_新規数  = find_col(ws, header_row_idx, ['新規数合計', '新規数'])
    col_指名率  = find_col(ws, header_row_idx, ['指名\n比率', '指名比率'])
    col_空き枠率 = find_col(ws, header_row_idx, ['空き枠率', '空き率', '空枠率'])

    # シフト時間と勤務時間が近接している場合、より右の列（勤務時間）を使う
    col_シフト = find_col(ws, header_row_idx, ['シフト時間', '勤務予定時間'])
    if col_稼働 and col_シフト and col_稼働 == col_シフト:
        # 同じ列になってしまった場合は右隣を勤務時間とする
        col_稼働 = col_シフト + 1

    # ── データ行を抽出 ───────────────────────────────────
    SKIP_VALUES = {'0', '', 'アルバイト', '体験', '面貸し'}
    STOP_KEYWORDS = ('①合計', '②面貸し', '総売上', '面貸し合計')

    stylists = []
    data_start = header_row_idx + 2  # ヘッダー2行分 + 1

    for row in ws.iter_rows(min_row=data_start, values_only=False):
        name_cell = row[stylist_col - 1]  # 0-indexed
        name_val = name_cell.value

        if name_val is None:
            continue

        name_str = str(name_val).strip()

        # 終了キーワードに達したらストップ
        if any(kw in name_str for kw in STOP_KEYWORDS):
            break

        # スキップ対象
        if name_str in SKIP_VALUES or name_str.isdigit():
            continue

        def get(col):
            if col is None:
                return None
            return clean_number(row[col - 1].value)

        stylist = {
            'name':    name_str,
            '生産性':  get(col_生産性),
            '売上':    get(col_売上),
            '稼働時間': get(col_稼働),
            '客単価':  get(col_客単価),
            '新規数':  get(col_新規数),
            '指名率':  get(col_指名率),
            '空き枠率': get(col_空き枠率),
        }

        # 売上が0またはNoneのスタイリストはスキップ
        if not stylist['売上']:
            continue

        stylists.append(stylist)

    return {
        'store':    store_name,
        'stylists': stylists,
    }
